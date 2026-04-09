"""
Excel 导出服务
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
from typing import List

from app.models.schedule import ClassSession


def generate_schedule_excel(sessions: List[ClassSession]) -> bytes:
    """生成课表 Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "课程表"

    # 表头样式
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 表头
    headers = ['日期', '时间', '课程', '教师', '教室', '学生数', '状态', '备注']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 设置列宽
    column_widths = [12, 15, 20, 10, 15, 8, 10, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 数据行
    for row, session in enumerate(sessions, 2):
        time_slot = session.time_slot
        data = [
            str(session.session_date),
            f"{time_slot.start_time}-{time_slot.end_time}" if time_slot else "",
            session.course.name if session.course else "",
            session.teacher.name if session.teacher else "",
            session.classroom.name if session.classroom else "",
            len(session.student_attendances) if session.student_attendances else 0,
            session.status.value,
            session.notes or ""
        ]

        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = center_align
            cell.border = thin_border

    # 保存到字节流
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()